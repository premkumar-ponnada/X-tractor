"""Generate a test pack with one file per supported format (synthetic Swedish tender content).

    python scripts/make_format_pack.py            # → X-tractor/test-files/format-pack/  (+ format-pack.zip)

Upload the whole folder (or the zip) in X-tractor to run every SDK on every format.
Legacy/OpenDocument formats are produced with LibreOffice; skipped with a note if it is missing.
Outlook .msg cannot be generated without Outlook, so it is not included.
"""

import shutil
import subprocess
import sys
import tempfile
import zipfile
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "test-files" / "format-pack"
sys.path.insert(0, str(ROOT / "backend"))
sys.stdout.reconfigure(encoding="utf-8")

from extractors.libreoffice import ensure_on_path  # noqa: E402
from tests.make_fixtures import make_docx, make_scanned_pdf, make_text_pdf, make_xlsx  # noqa: E402

TITLE = "Förfrågningsunderlag – Rivning Ödla 4"
PARAGRAPHS = [
    "Entreprenaden omfattar rivning av två byggnader samt sanering av asbest och PCB.",
    "Anbud ska lämnas senast 2026-11-02 via upphandlingsportalen. Anbudet ska vara giltigt i 90 dagar.",
    "Arbetstid: vardagar 07.00–16.00. Platsbesök sker 2026-10-14 kl. 10.00.",
]
ROWS = [
    ("Pos", "Beskrivning", "Mängd", "Enhet"),
    ("1", "Rivning av tak", "120", "m²"),
    ("2", "Asbestsanering", "45", "m²"),
    ("3", "Bortforsling", "80", "ton"),
]


def write(name: str, content: str) -> None:
    (OUT / name).write_text(content, encoding="utf-8")


def make_text_formats() -> None:
    table_md = (
        "\n".join("| " + " | ".join(r) + " |" for r in ROWS[:1])
        + "\n|---|---|---|---|\n"
        + "\n".join("| " + " | ".join(r) + " |" for r in ROWS[1:])
    )
    write("beskrivning.md", f"# {TITLE}\n\n" + "\n\n".join(PARAGRAPHS) + f"\n\n## Mängder\n\n{table_md}\n")
    write("beskrivning.txt", f"{TITLE}\n\n" + "\n\n".join(PARAGRAPHS) + "\n")
    rows_html = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in ROWS)
    write(
        "beskrivning.html",
        f"<!doctype html><html lang='sv'><head><meta charset='utf-8'><title>{TITLE}</title></head><body>"
        f"<h1>{TITLE}</h1>" + "".join(f"<p>{p}</p>" for p in PARAGRAPHS) + f"<table>{rows_html}</table></body></html>",
    )
    items = "".join(
        f"<post nr='{r[0]}'><beskrivning>{r[1]}</beskrivning><mangd enhet='{r[3]}'>{r[2]}</mangd></post>" for r in ROWS[1:]
    )
    write("mangder.xml", f"<?xml version='1.0' encoding='UTF-8'?><forfragan titel='{TITLE}'>{items}</forfragan>")
    write("mangder.csv", "\n".join(";".join(r) for r in ROWS) + "\n")


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = TITLE
    slide.placeholders[1].text = PARAGRAPHS[0]
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Mängdförteckning"
    table = slide.shapes.add_table(len(ROWS), 4, Inches(0.5), Inches(1.5), Inches(9), Inches(2)).table
    for r, row in enumerate(ROWS):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    deck.save(path)


def make_images() -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default()
    for i, line in enumerate([TITLE, *PARAGRAPHS[:2]]):
        draw.text((60, 80 + i * 110), line[:70], fill="black", font=font)
    image.save(OUT / "skannad-sida.png")
    image.convert("L").save(OUT / "skannad-sida.jpg", quality=90)
    image.convert("L").save(OUT / "skannad-sida.tiff")


def make_eml(path: Path, attachment: Path) -> None:
    message = EmailMessage()
    message["From"] = "upphandling@kommunen.se"
    message["To"] = "anbud@daco.se"
    message["Subject"] = f"Förfrågan: {TITLE}"
    message["Date"] = "Mon, 05 Oct 2026 09:00:00 +0200"
    message.set_content("Hej!\n\n" + "\n".join(PARAGRAPHS) + "\n\nMvh Kommunen")
    message.add_alternative("<p>Hej!</p>" + "".join(f"<p>{p}</p>" for p in PARAGRAPHS), subtype="html")
    message.add_attachment(attachment.read_bytes(), maintype="application", subtype="pdf", filename=attachment.name)
    path.write_bytes(message.as_bytes())


def make_epub(path: Path) -> None:
    body = "".join(f"<p>{p}</p>" for p in PARAGRAPHS)
    with zipfile.ZipFile(path, "w") as epub:
        epub.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        epub.writestr(
            "META-INF/container.xml",
            "<?xml version='1.0'?><container version='1.0' xmlns='urn:oasis:names:tc:opendocument:xmlns:container'>"
            "<rootfiles><rootfile full-path='OEBPS/content.opf' media-type='application/oebps-package+xml'/></rootfiles></container>",
        )
        epub.writestr(
            "OEBPS/content.opf",
            "<?xml version='1.0' encoding='UTF-8'?><package xmlns='http://www.idpf.org/2007/opf' version='3.0' unique-identifier='id'>"
            f"<metadata xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:identifier id='id'>xtractor-test</dc:identifier><dc:title>{TITLE}</dc:title><dc:language>sv</dc:language></metadata>"
            "<manifest><item id='c1' href='kapitel.xhtml' media-type='application/xhtml+xml'/><item id='nav' href='nav.xhtml' media-type='application/xhtml+xml' properties='nav'/></manifest>"
            "<spine><itemref idref='c1'/></spine></package>",
        )
        epub.writestr(
            "OEBPS/nav.xhtml",
            "<?xml version='1.0' encoding='UTF-8'?><html xmlns='http://www.w3.org/1999/xhtml' xmlns:epub='http://www.idpf.org/2007/ops'><body>"
            "<nav epub:type='toc'><ol><li><a href='kapitel.xhtml'>Kapitel</a></li></ol></nav></body></html>",
        )
        epub.writestr(
            "OEBPS/kapitel.xhtml",
            f"<?xml version='1.0' encoding='UTF-8'?><html xmlns='http://www.w3.org/1999/xhtml'><head><title>{TITLE}</title></head><body><h1>{TITLE}</h1>{body}</body></html>",
        )


def convert(source: Path, target_ext: str, name: str) -> bool:
    """LibreOffice headless conversion into the pack (legacy + OpenDocument formats)."""
    with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as profile:
        result = subprocess.run(
            [
                shutil.which("soffice"),
                f"-env:UserInstallation={Path(profile).as_uri()}",
                "--headless",
                "--convert-to",
                target_ext,
                "--outdir",
                tmp,
                str(source),
            ],
            capture_output=True,
            timeout=180,
        )
        produced = Path(tmp) / f"{source.stem}.{target_ext.split(':')[0]}"
        if result.returncode == 0 and produced.exists():
            shutil.move(produced, OUT / name)
            return True
    return False


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    make_text_pdf(OUT / "af-text.pdf")
    make_scanned_pdf(OUT / "miljoinventering-skannad.pdf")
    make_docx(OUT / "teknisk-beskrivning.docx")
    make_xlsx(OUT / "mangdforteckning.xlsx")
    make_pptx(OUT / "presentation.pptx")
    make_text_formats()
    make_images()
    make_eml(OUT / "forfragan.eml", OUT / "af-text.pdf")
    make_epub(OUT / "handbok.epub")

    if ensure_on_path():
        for source, ext, name in [
            ("teknisk-beskrivning.docx", "doc", "teknisk-beskrivning.doc"),
            ("teknisk-beskrivning.docx", "rtf", "teknisk-beskrivning.rtf"),
            ("teknisk-beskrivning.docx", "odt", "teknisk-beskrivning.odt"),
            ("mangdforteckning.xlsx", "xls", "mangdforteckning.xls"),
            ("mangdforteckning.xlsx", "ods", "mangdforteckning.ods"),
            ("presentation.pptx", "ppt", "presentation.ppt"),
            ("presentation.pptx", "odp", "presentation.odp"),
        ]:
            print(("  ✓ " if convert(OUT / source, ext, name) else "  ✗ ") + name)
    else:
        print("LibreOffice not found — legacy/OpenDocument files skipped")

    # A nested archive like real tender packages: zip containing a folder and another zip.
    inner = OUT.parent / "bilagor.zip"
    with zipfile.ZipFile(inner, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(OUT / "mangdforteckning.xlsx", "bilagor/mangdforteckning.xlsx")
        archive.write(OUT / "beskrivning.html", "bilagor/beskrivning.html")
    with zipfile.ZipFile(OUT / "anbudspaket.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(OUT / "af-text.pdf", "Förfrågningsunderlag/AF.pdf")
        archive.write(OUT / "teknisk-beskrivning.docx", "Förfrågningsunderlag/Teknisk beskrivning.docx")
        archive.write(inner, "Förfrågningsunderlag/bilagor.zip")
    inner.unlink()

    files = sorted(p.name for p in OUT.iterdir())
    bundle = OUT.parent / "format-pack.zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in files:
            archive.write(OUT / name, name)
    print(f"\n{len(files)} files in {OUT}:")
    print("  " + ", ".join(files))
    print(f"Bundle: {bundle}")


if __name__ == "__main__":
    main()
