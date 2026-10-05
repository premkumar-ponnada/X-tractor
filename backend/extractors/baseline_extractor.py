"""Baseline: simple per-format readers (pypdf, python-docx, openpyxl, xlrd, extract-msg).

Mirrors what a typical hand-built pipeline (like DACO's current one) does, so every other
SDK has a reference point. Text layer only — no OCR, no layout analysis.
"""

import email
import json
import re
from email import policy
from html import unescape
from pathlib import Path

from extractors.base import (
    EmitFn,
    ExtractionError,
    ExtractionResult,
    ExtractOptions,
    FileKind,
    OutputFormat,
    Page,
    package_version,
)

KINDS = {
    FileKind.PDF,
    FileKind.DOCX,
    FileKind.XLSX,
    FileKind.XLS,
    FileKind.TXT,
    FileKind.CSV,
    FileKind.MD,
    FileKind.EML,
    FileKind.MSG,
    FileKind.HTML,
}
TAG = re.compile(r"<[^>]+>")
SPACE = re.compile(r"[ \t]+")


def decode_text(raw: bytes) -> str:
    for codec in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(codec)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</h\d>", "\n", html)
    return SPACE.sub(" ", unescape(TAG.sub(" ", html))).strip()


class BaselineExtractor:
    name = "baseline"

    def version(self) -> str | None:
        return f"pypdf {package_version('pypdf')}"

    def supports(self, kind: FileKind) -> bool:
        return kind in KINDS

    def output_formats(self) -> list[OutputFormat]:
        return [OutputFormat.TEXT, OutputFormat.JSON]

    def availability(self) -> tuple[bool, str]:
        return True, "Built-in readers"

    def extract(self, path: Path, kind: FileKind, opts: ExtractOptions, emit: EmitFn) -> ExtractionResult:
        warnings: list[str] = []
        reader = {
            FileKind.PDF: self._pdf,
            FileKind.DOCX: self._docx,
            FileKind.XLSX: self._xlsx,
            FileKind.XLS: self._xls,
            FileKind.EML: self._eml,
            FileKind.MSG: self._msg,
            FileKind.HTML: lambda p, w, e: [Page(n=1, text=html_to_text(decode_text(p.read_bytes())))],
        }.get(kind, lambda p, w, e: [Page(n=1, text=decode_text(p.read_bytes()))])
        emit("parse", f"Reading {kind.value.upper()} with built-in reader", None)
        pages = reader(path, warnings, emit)
        text = "\n\n".join(f"--- Page {p.n} ---\n{p.text}" for p in pages)
        return ExtractionResult(
            pages=pages,
            outputs={
                OutputFormat.TEXT: text,
                OutputFormat.JSON: json.dumps([p.model_dump() for p in pages], ensure_ascii=False, indent=2),
            },
            warnings=warnings,
        )

    def _pdf(self, path: Path, warnings: list[str], emit: EmitFn) -> list[Page]:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(path)
            if reader.is_encrypted:
                reader.decrypt("")
        except (PdfReadError, ValueError) as exc:
            raise ExtractionError(f"Could not open PDF: {exc}") from exc
        pages: list[Page] = []
        total = len(reader.pages)
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:  # one broken page must not lose the document
                warnings.append(f"Page {i}: {exc}")
                text = ""
            try:
                images = len(page.images)
            except Exception:
                images = 0
            pages.append(Page(n=i, text=text, images=images))
            if i % 10 == 0 or i == total:
                emit("page", f"Page {i}/{total}", {"page": i, "total": total})
        return pages

    def _docx(self, path: Path, warnings: list[str], emit: EmitFn) -> list[Page]:
        import docx

        document = docx.Document(str(path))
        parts: list[str] = []
        for section in document.sections:
            header = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
            if header:
                parts.append(header)
        parts.extend(p.text for p in document.paragraphs if p.text.strip())
        for table in document.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            parts.append("\n".join(rows))
        for section in document.sections:
            footer = "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())
            if footer:
                parts.append(footer)
        warnings.append("Word files have no fixed pages — returned as one page")
        return [Page(n=1, text="\n\n".join(parts), tables=len(document.tables), images=len(document.inline_shapes))]

    def _xlsx(self, path: Path, warnings: list[str], emit: EmitFn) -> list[Page]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        pages = []
        for n, sheet in enumerate(workbook.worksheets, start=1):
            rows = ["\t".join("" if c is None else str(c) for c in row) for row in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if r.strip()]
            pages.append(Page(n=n, text=f"# {sheet.title}\n" + "\n".join(rows), tables=1 if rows else 0))
        workbook.close()
        return pages

    def _xls(self, path: Path, warnings: list[str], emit: EmitFn) -> list[Page]:
        import xlrd

        book = xlrd.open_workbook(str(path))
        pages = []
        for n, sheet in enumerate(book.sheets(), start=1):
            rows = ["\t".join(str(v) for v in sheet.row_values(r)) for r in range(sheet.nrows)]
            pages.append(Page(n=n, text=f"# {sheet.name}\n" + "\n".join(rows), tables=1 if rows else 0))
        return pages

    def _eml(self, path: Path, warnings: list[str], emit: EmitFn) -> list[Page]:
        message = email.message_from_bytes(path.read_bytes(), policy=policy.default)
        head = "\n".join(f"{h}: {message.get(h, '')}" for h in ("From", "To", "Date", "Subject"))
        body_part = message.get_body(preferencelist=("plain", "html"))
        body = ""
        if body_part is not None:
            body = body_part.get_content()
            if body_part.get_content_type() == "text/html":
                body = html_to_text(body)
        attachments = [a.get_filename() or "unnamed" for a in message.iter_attachments()]
        if attachments:
            warnings.append(f"Attachments not read by baseline: {', '.join(attachments)}")
        return [Page(n=1, text=f"{head}\n\n{body}")]

    def _msg(self, path: Path, warnings: list[str], emit: EmitFn) -> list[Page]:
        import extract_msg

        with extract_msg.openMsg(str(path)) as message:
            head = f"From: {message.sender}\nTo: {message.to}\nDate: {message.date}\nSubject: {message.subject}"
            body = message.body or ""
            attachments = [getattr(a, "longFilename", None) or "unnamed" for a in message.attachments]
        if attachments:
            warnings.append(f"Attachments not read by baseline: {', '.join(attachments)}")
        return [Page(n=1, text=f"{head}\n\n{body}")]
