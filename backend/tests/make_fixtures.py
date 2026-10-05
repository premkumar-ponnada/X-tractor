"""Generate small synthetic test files (no client data): python tests/make_fixtures.py

Covers: multi-page text PDF with Swedish characters, scanned (image-only) PDF, DOCX with
header/footer/table, XLSX with two sheets, CSV, EML with an attachment, and a ZIP of them.
"""

import zipfile
from email.message import EmailMessage
from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures"

PAGES = [
    [
        "Administrativa föreskrifter (AF)",
        "AFB.31 Tider: Entreprenaden ska vara färdigställd 2026-12-18.",
        "AFC.11 Ersättningsform: Fast pris. Rivning av två byggnader på fastigheten Ödla 4.",
    ],
    [
        "Mängdförteckning",
        "Pos 1  Rivning av tak           120 m2",
        "Pos 2  Sanering av asbest        45 m2",
        "Pos 3  Bortforsling av massor    80 ton",
    ],
]


def _pdf_text(lines: list[str]) -> bytes:
    ops = ["BT", "/F1 14 Tf", "72 760 Td", "18 TL"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({escaped}) Tj T*")
    ops.append("ET")
    return "\n".join(ops).encode("cp1252")


def make_text_pdf(path: Path) -> None:
    """Hand-built PDF: Helvetica with WinAnsiEncoding covers å, ä, ö."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(PAGES))]
    font_id = 3 + 2 * len(PAGES)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(PAGES)} >>".encode())
    for i, lines in enumerate(PAGES):
        content = _pdf_text(lines)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {page_ids[i] + 1} 0 R >>".encode()
        )
        objects.append(b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")

    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def make_scanned_pdf(path: Path) -> None:
    """Text rendered as an image only — no text layer, so only OCR can read it."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 40)
    except OSError:
        font = ImageFont.load_default()
    for i, line in enumerate(["Miljöinventering", "Asbest påträffad i fasadplattor.", "Provtagning utförd 2026-09-01."]):
        draw.text((100, 150 + i * 80), line, fill="black", font=font)
    image.save(path, "PDF", resolution=150)


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "DACO Contractor – Förfrågningsunderlag"
    document.sections[0].footer.paragraphs[0].text = "Sida 1 – Konfidentiellt"
    document.add_heading("Teknisk beskrivning", level=1)
    document.add_paragraph("Rivningsarbeten ska utföras enligt AMA Anläggning 23.")
    table = document.add_table(rows=3, cols=3)
    for r, row in enumerate([("Pos", "Arbete", "Mängd"), ("1", "Rivning", "120 m²"), ("2", "Sanering", "45 m²")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.save(path)


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mängder"
    for row in [("Pos", "Beskrivning", "Mängd", "Enhet"), (1, "Rivning tak", 120, "m2"), (2, "Asbestsanering", 45, "m2")]:
        sheet.append(row)
    prices = workbook.create_sheet("Priser")
    prices.append(("Pos", "À-pris (kr)"))
    prices.append((1, 350))
    workbook.save(path)


def make_csv(path: Path) -> None:
    path.write_text("pos;beskrivning;mängd\n1;Rivning av förråd;30\n2;Håltagning;12\n", encoding="utf-8")


def make_eml(path: Path, attachment: Path) -> None:
    message = EmailMessage()
    message["From"] = "upphandling@kommun.se"
    message["To"] = "anbud@daco.se"
    message["Subject"] = "Förfrågan: Rivning Ödla 4"
    message["Date"] = "Mon, 05 Oct 2026 09:00:00 +0200"
    message.set_content("Hej!\n\nBifogat finns förfrågningsunderlaget. Anbud lämnas senast 2026-11-02.\n\nMvh Kommunen")
    message.add_attachment(attachment.read_bytes(), maintype="application", subtype="pdf", filename=attachment.name)
    path.write_bytes(message.as_bytes())


def main() -> None:
    FIXTURES.mkdir(exist_ok=True)
    text_pdf = FIXTURES / "af-text.pdf"
    make_text_pdf(text_pdf)
    make_scanned_pdf(FIXTURES / "miljo-scanned.pdf")
    make_docx(FIXTURES / "teknisk-beskrivning.docx")
    make_xlsx(FIXTURES / "mangder.xlsx")
    make_csv(FIXTURES / "poster.csv")
    make_eml(FIXTURES / "forfragan.eml", text_pdf)
    with zipfile.ZipFile(FIXTURES / "anbud-paket.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for name in ("af-text.pdf", "teknisk-beskrivning.docx", "mangder.xlsx"):
            archive.write(FIXTURES / name, f"underlag/{name}")
        archive.writestr("__MACOSX/._junk", b"junk")
    print("\n".join(sorted(p.name for p in FIXTURES.iterdir())))


if __name__ == "__main__":
    main()
